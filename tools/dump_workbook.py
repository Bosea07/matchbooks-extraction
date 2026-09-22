"""Print the shape and head of every sheet in a workbook.

    py -B tools\\dump_workbook.py "C:\\path\\to\\book.xlsx" [rows]

Used to see what a multi-tab reconciliation actually contains before pointing
the engine at it. Default 25 rows per sheet; pass a number to change it.
"""
import os
import sys

MAX_CELL = 30


def cell(v):
    if v is None:
        return ''
    s = str(v).replace('\n', '\\n').replace('\r', '')
    return s[:MAX_CELL]


def main(path, nrows):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    print('=' * 78)
    print('FILE  ', os.path.basename(path))
    print('SHEETS', wb.sheetnames)
    print('=' * 78)

    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        # trim wholly empty trailing rows
        while rows and not any(c not in (None, '') for c in rows[-1]):
            rows.pop()
        width = max((len(r) for r in rows), default=0)
        print(f'\n{"#" * 74}')
        print(f'# SHEET {ws.title!r} — {len(rows)} rows x {width} cols')
        print(f'{"#" * 74}')
        for i, r in enumerate(rows[:nrows]):
            line = ' | '.join(cell(c) for c in r)
            print(f'{i:>4}| {line[:230]}')
        if len(rows) > nrows:
            print(f'     ... {len(rows) - nrows} more rows')
            print(f'     LAST ROW: ' +
                  ' | '.join(cell(c) for c in rows[-1])[:230])

        # numeric column totals help identify which column is the amount
        sums, counts = {}, {}
        for r in rows:
            for j, c in enumerate(r):
                if isinstance(c, (int, float)):
                    sums[j] = sums.get(j, 0) + c
                    counts[j] = counts.get(j, 0) + 1
        if sums:
            print('     numeric columns (index: count, sum):')
            for j in sorted(sums):
                print(f'        col {j:>2}: {counts[j]:>4} values, {sums[j]:>18,.2f}')


if __name__ == '__main__':
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    main(sys.argv[1], int(sys.argv[2]) if len(sys.argv) > 2 else 25)
